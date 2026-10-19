import re
from copy import deepcopy
from docx import Document
from docx.table import Table

DOCX = r"C:\Users\galazel\IdeaProjects\capstone\frontend\REBYU_Chapter_III_Revised_Formatted.docx"
SOURCE = r"C:\Users\galazel\IdeaProjects\capstone\DATA_DICTIONARY.md"

def parse_dictionary(markdown):
    entries = []
    pattern = re.compile(
        r"### DATA DICTIONARY OF ([A-Z0-9_]+)\r?\n\r?\n"
        r"\| Field Name \| Constraints \| Data Type \| Allow Nulls \| Description \|\r?\n"
        r"\|---\|---\|---\|---\|---\|\r?\n"
        r"((?:\|.*\|\r?\n)+)",
    )
    for match in pattern.finditer(markdown):
        rows = []
        for line in match.group(2).strip().splitlines():
            values = [value.strip() for value in line.strip().strip("|").split("|")]
            if len(values) == 5:
                rows.append(values)
        entries.append((match.group(1), rows))
    return entries

def clear_cell(cell, text):
    cell.text = text

def cloned_table(document, template, rows):
    table_xml = deepcopy(template._tbl)
    existing_rows = list(table_xml.tr_lst)
    header = existing_rows[0]
    sample = existing_rows[1]
    for row in existing_rows[1:]:
        table_xml.remove(row)
    for _ in rows:
        table_xml.append(deepcopy(sample))
    table = Table(table_xml, document._body)
    for cell, value in zip(table.rows[0].cells, ["Field Name", "Constraints", "Data Type", "Allow Nulls", "Description"]):
        clear_cell(cell, value)
    for row, values in zip(table.rows[1:], rows):
        for cell, value in zip(row.cells, values):
            clear_cell(cell, value)
    return table_xml

with open(SOURCE, encoding="utf-8") as file:
    entries = parse_dictionary(file.read())
if len(entries) != 102:
    raise RuntimeError(f"Expected 102 live-schema entries; found {len(entries)}")

document = Document(DOCX)
paragraphs = document.paragraphs
start = next(p for p in paragraphs if p.text.strip() == "DATA DICTIONARY OF USER_TYPES")
network = next(p for p in paragraphs if p.text.strip() == "Network Design")
template = next(t for t in document.tables if t.rows and t.rows[0].cells[0].text.strip() == "Field Name")
body = document.element.body

# Remove the old, partially inferred dictionary in one contiguous body range.
start_pos = list(body).index(start._p)
end_pos = list(body).index(network._p)
for child in list(body)[start_pos:end_pos]:
    body.remove(child)

# Create exact table entries from DATA_DICTIONARY.md, whose source is the live
# Supabase public schema.  The original manuscript table grid is cloned.
new_elements = []
for number, (name, rows) in enumerate(entries, start=6):
    title = document.add_paragraph(f"DATA DICTIONARY OF {name}")
    title.style = "Normal"
    caption = document.add_paragraph(f"Table {number}: Data Dictionary of {name.replace('_', ' ').title()}")
    caption.style = "Table Caption"
    new_elements.extend([title._p, caption._p])
    table_xml = cloned_table(document, template, rows)
    body.append(table_xml)
    new_elements.append(table_xml)

# The newly built blocks were appended temporarily; place them in their
# original location immediately before Network Design.
for element in new_elements:
    network._p.addprevious(element)

# Continue numbering after the 102 live-schema dictionary tables.
replacements = {
    "Table 105: Technology Stack of REBYU": "Table 108: Technology Stack of REBYU",
    "Table 106: Software Specification": "Table 109: Software Specification",
    "Table 107: Hardware Specification": "Table 110: Hardware Specification",
    "Table 108: List of Modules": "Table 111: List of Modules",
}
for paragraph in document.paragraphs:
    if paragraph.text in replacements:
        paragraph.text = replacements[paragraph.text]
        paragraph.style = "Table Caption"

document.save(DOCX)
