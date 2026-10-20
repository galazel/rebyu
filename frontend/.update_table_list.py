from copy import deepcopy
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = r"C:\Users\galazel\IdeaProjects\capstone\frontend\REBYU_Chapter_III_Revised_Formatted.docx"
document = Document(SOURCE)

# A dedicated caption style lets Word/WPS maintain the List of Tables from the
# current manuscript captions rather than the pre-revision list.
styles = document.styles
try:
    caption_style = styles["Table Caption"]
except KeyError:
    caption_style = styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption_style.base_style = styles["Normal"]
    caption_style.font.name = "Times New Roman"
    caption_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

def set_caption(paragraph, text):
    paragraph.text = text
    paragraph.style = caption_style

# Existing non-dictionary tables, using sequential unique numbering.
fixed = {
    310: "Table 2: Distribution of Respondents",
    329: "Table 3: Scoring Points",
    334: "Table 4: REBYU's Ethical Principle",
    806: "Table 105: Technology Stack of REBYU",
    810: "Table 106: Software Specification",
    813: "Table 107: Hardware Specification",
    846: "Table 108: List of Modules",
}

# Table 1 is the comparative matrix immediately preceding the related-system
# discussion; it did not have a standalone visible caption in the template.
set_caption(document.paragraphs[281], "Table 1: Comparative Matrix of Related Systems")
for index, caption in fixed.items():
    set_caption(document.paragraphs[index], caption)

# Data dictionary captions: only entities represented in the current codebase.
dictionary_number = 6
for index, paragraph in enumerate(document.paragraphs[:-1]):
    title = paragraph.text.strip()
    if not title.startswith("DATA DICTIONARY OF "):
        continue
    next_paragraph = document.paragraphs[index + 1]
    if not next_paragraph.text.strip().startswith("Table "):
        continue
    name = title.removeprefix("DATA DICTIONARY OF ").replace("_", " ").title()
    set_caption(next_paragraph, f"Table {dictionary_number}: Data Dictionary of {name}")
    dictionary_number += 1

if dictionary_number != 105:
    raise RuntimeError(f"Expected 99 data-dictionary tables; found {dictionary_number - 6}")

# Replace the old blank List-of-Tables placeholder with a live Table of Figures
# field filtered to the Table Caption style.  Word/WPS refreshes it on open.
lot_heading = next(p for p in document.paragraphs if p.text.strip().upper() == "LIST OF TABLES")
figure_heading = next(p for p in document.paragraphs if p.text.strip().upper() == "LIST OF FIGURES")
body = figure_heading._p.getparent()
start = list(body).index(lot_heading._p) + 1
end = list(body).index(figure_heading._p)
for element in list(body)[start:end]:
    body.remove(element)

field = OxmlElement("w:p")
p_pr = OxmlElement("w:pPr")
field.append(p_pr)
for kind, content in (("begin", None), (None, ' TOC \\h \\z \\t "Table Caption,1" '), ("separate", None)):
    run = OxmlElement("w:r")
    if kind:
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), kind)
        run.append(char)
    else:
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = content
        run.append(instr)
    field.append(run)
result = OxmlElement("w:r")
text = OxmlElement("w:t")
text.text = "Update this List of Tables field to display the revised table entries."
result.append(text)
field.append(result)
end_run = OxmlElement("w:r")
end_char = OxmlElement("w:fldChar")
end_char.set(qn("w:fldCharType"), "end")
end_run.append(end_char)
field.append(end_run)
figure_heading._p.addprevious(field)

# Request refresh on opening in Word/WPS.
settings = document.settings.element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

document.save(SOURCE)
